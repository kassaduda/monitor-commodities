"""Leitura do Pink Sheet do Banco Mundial a partir do arquivo baixado manualmente.

Não há requisição de rede: o link do arquivo muda a cada edição, então a usuária baixa
CMO-Historical-Data-Monthly.xlsx e o coloca em data/raw/worldbank/.

Layout conferido em 14/10/2026 (edição "Updated on October 02, 2026"): aba "Monthly Prices",
linhas de título (uma delas "Updated on ..."), uma linha com os nomes das séries, a linha
seguinte com as unidades e uma linha por mês com o período na coluna A ("AAAAMmm").
"""

import re
from dataclasses import dataclass
from datetime import date, datetime
from pathlib import Path

from openpyxl import load_workbook

from monitor import config

PADRAO_PERIODO = re.compile(r"^\d{4}M(0[1-9]|1[0-2])$")
PADRAO_ATUALIZADO = re.compile(r"Updated on (\w+ \d{1,2}, \d{4})")
MAX_LINHAS_CABECALHO = 15


class ErroPinkSheet(ValueError):
    """Arquivo do Pink Sheet fora do layout esperado."""


@dataclass(frozen=True)
class SeriePinkSheet:
    nome: str
    unidade: str
    registros: list[tuple[str, object]]  # (período "AAAAMmm", valor bruto da célula)


@dataclass(frozen=True)
class PinkSheet:
    atualizado_em: date | None
    series: dict[str, SeriePinkSheet]


def caminho_arquivo(dir_raw: Path = config.DIR_RAW) -> Path:
    return dir_raw / "worldbank" / config.WB_ARQUIVO


def _data_atualizacao(texto: object) -> date | None:
    m = PADRAO_ATUALIZADO.search(str(texto or ""))
    return datetime.strptime(m.group(1), "%B %d, %Y").date() if m else None


def _indices_das_colunas(cabecalho: tuple, colunas: list[str]) -> dict[str, int] | None:
    nomes = [str(c).strip() if c is not None else "" for c in cabecalho]
    if not all(c in nomes for c in colunas):
        return None
    repetidas = [c for c in colunas if nomes.count(c) > 1]
    if repetidas:
        raise ErroPinkSheet(f"Colunas repetidas no cabeçalho: {repetidas}")
    return {c: nomes.index(c) for c in colunas}


def ler_pink_sheet(caminho: Path, colunas: list[str], aba: str = config.WB_ABA) -> PinkSheet:
    """Lê as colunas pedidas da aba de preços mensais, sem converter os valores."""
    wb = load_workbook(caminho, read_only=True, data_only=True)
    try:
        if aba not in wb.sheetnames:
            raise ErroPinkSheet(f"Aba {aba!r} não encontrada; abas: {wb.sheetnames}")
        linhas = wb[aba].iter_rows(values_only=True)

        atualizado_em: date | None = None
        indices: dict[str, int] | None = None
        for _ in range(MAX_LINHAS_CABECALHO):
            linha = next(linhas, None)
            if linha is None:
                break
            atualizado_em = atualizado_em or _data_atualizacao(linha[0] if linha else None)
            indices = _indices_das_colunas(linha, colunas)
            if indices:
                break
        if not indices:
            raise ErroPinkSheet(f"Cabeçalho com {colunas} não encontrado nas primeiras linhas")

        unidades = next(linhas, ())
        registros: dict[str, list[tuple[str, object]]] = {c: [] for c in colunas}
        for linha in linhas:
            periodo = str(linha[0]).strip() if linha and linha[0] is not None else ""
            if not PADRAO_PERIODO.match(periodo):
                continue
            for c, j in indices.items():
                registros[c].append((periodo, linha[j] if j < len(linha) else None))
    finally:
        wb.close()

    series = {
        c: SeriePinkSheet(c, str(unidades[j]).strip() if j < len(unidades) else "", registros[c])
        for c, j in indices.items()
    }
    return PinkSheet(atualizado_em, series)
